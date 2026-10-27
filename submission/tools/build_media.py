#!/usr/bin/env python3
"""
Build the submission slide deck (.pptx) and the narrated demo video (.mp4).

Inputs : real screenshots of the running app (--frames) and a running API in SAMPLE_MODE for the terminal scene.
Outputs: <out>/AEGIS_Slide_Deck.pptx, <out>/AEGIS_Demo_Video.mp4, <out>/video_frames/*.png, <out>/narration.txt

Dependencies (a throwaway venv is enough): python-pptx, pillow, imageio-ffmpeg. Narration uses the macOS `say` voice.
Nothing here invents results: the terminal scene is rendered from a real POST /run response and its trace file,
and every number on a slide is quoted from the repo (see submission/*.md).

  python submission/tools/build_media.py --frames /path/to/screenshots --out submission/media [--api http://127.0.0.1:8000]
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
import textwrap
import urllib.request
import wave
from pathlib import Path

import imageio_ffmpeg
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
FFMPEG = imageio_ffmpeg.get_ffmpeg_exe()
W, H = 1920, 1080

BG, PANEL, PANEL2 = (7, 11, 20), (17, 24, 39), (26, 35, 54)
TEXT, MUTED = (230, 236, 245), (140, 152, 170)
TEAL, PURPLE, AMBER, RED, GREEN, BLUE = (45, 212, 191), (168, 85, 247), (245, 158, 11), (244, 63, 94), (16, 185, 129), (56, 189, 248)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype("/System/Library/Fonts/Avenir Next.ttc", size, index=0 if bold else 7)


def mono(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype("/System/Library/Fonts/Menlo.ttc", size, index=1 if bold else 0)


def canvas() -> tuple[Image.Image, ImageDraw.ImageDraw]:
    im = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, W, 8], fill=TEAL)
    return im, d


def wrap(d: ImageDraw.ImageDraw, text: str, f, width: int) -> list[str]:
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if d.textlength(t, font=f) <= width:
            cur = t
        else:
            lines.append(cur)
            cur = w
    return lines + ([cur] if cur else [])


def para(d, xy, text, f, fill=TEXT, width=1400, gap=10):
    x, y = xy
    for ln in wrap(d, text, f, width):
        d.text((x, y), ln, font=f, fill=fill)
        y += f.size + gap
    return y


def header(d, title: str, sub: str = ""):
    d.text((100, 70), title, font=font(64, True), fill=TEXT)
    if sub:
        d.text((100, 150), sub, font=font(30), fill=MUTED)


def box(d, xy, fill=PANEL, outline=PANEL2, r=18):
    d.rounded_rectangle(xy, radius=r, fill=fill, outline=outline, width=2)


def footer(d, text: str):
    d.text((100, H - 60), text, font=font(24), fill=MUTED)


# ───────────────────────── content shared by the deck and the video ─────────────────────────
def load_api(api: str, example: int = 2, container: str | None = None) -> dict:
    body = json.loads((ROOT / f"input_examples/example_{example}.json").read_text())
    body["options"] = {"sample_mode": True}
    req = urllib.request.Request(f"{api}/run", data=json.dumps(body).encode(), headers={"content-type": "application/json"})
    resp = json.loads(urllib.request.urlopen(req, timeout=300).read())
    if container:     # the trace file lives inside the container (the judge's `docker exec ... cat logs/trace-*.jsonl`)
        raw = subprocess.run(["docker", "exec", container, "cat", resp["log_file"]], capture_output=True, text=True, check=True).stdout
    else:
        raw = (ROOT / resp["log_file"]).read_text()
    trace = [json.loads(l) for l in raw.splitlines()]
    health = json.loads(urllib.request.urlopen(f"{api}/health", timeout=10).read())
    return {"resp": resp, "trace": trace, "health": health}


def shot(frames: Path, name: str) -> Image.Image | None:
    p = frames / f"{name}.png"
    return Image.open(p).convert("RGB") if p.exists() else None


def screen_frame(title: str, caption: str, img: Image.Image | None, crop: tuple | None = None) -> Image.Image:
    im, d = canvas()
    d.text((100, 40), title, font=font(46, True), fill=TEXT)
    if img is not None:
        if crop:
            img = img.crop(crop)
        scale = min((W - 200) / img.width, (H - 250) / img.height)
        sh = img.resize((int(img.width * scale), int(img.height * scale)), Image.LANCZOS)
        x, y = (W - sh.width) // 2, 120
        box(d, [x - 6, y - 6, x + sh.width + 6, y + sh.height + 6], fill=PANEL2, outline=TEAL, r=10)
        im.paste(sh, (x, y))
    para(d, (100, H - 100), caption, font(30), MUTED, width=1700)
    return im


def terminal_frame(api_data: dict) -> Image.Image:
    resp, trace = api_data["resp"], api_data["trace"]
    im, d = canvas()
    d.text((100, 36), "The judge path: one command, no credentials", font=font(46, True), fill=TEXT)
    box(d, [100, 110, W - 100, H - 40], fill=(10, 14, 24), outline=PANEL2)
    d.rectangle([100, 110, W - 100, 146], fill=PANEL2)
    for i, c in enumerate((RED, AMBER, GREEN)):
        d.ellipse([122 + i * 30, 120, 138 + i * 30, 136], fill=c)
    f, fb = mono(22), mono(22, True)
    y, step = 160, 29
    maxc = 124

    def line(text, col=TEXT, bold=False):
        nonlocal y
        d.text((124, y), text[:maxc], font=fb if bold else f, fill=col)
        y += step

    short = lambda m: (m or "none").replace("nvidia/", "").replace("nemotron-", "")
    line("$ docker build -t oah-submission . && docker run -p 8000:8000 -e SAMPLE_MODE=true oah-submission", GREEN, True)
    line("$ curl -X POST localhost:8000/run -H 'Content-Type: application/json' -d @input_examples/example_2.json", GREEN, True)
    o = resp["output"]
    line(f"  status={resp['status']}  stopping_reason={o['stopping_reason']}  data_mode={o['data_mode']}  links={len(o['links'])}")
    line("  agents: " + "  ".join(f"{a['agent']} x{a['calls']}" for a in resp["agents"]))
    line("  answered by: " + ", ".join(sorted({short(m) for a in resp["agents"] for m in a["reported_models"]})), BLUE)
    line("  engines: " + "  ".join(f"{e['engine']}={e['status']}" for e in resp["engines"]), AMBER)
    y += 6
    line(f"$ docker exec <container> cat {resp['log_file']}", GREEN, True)
    shown, last_handoff = 0, False
    for e in trace:
        pm, sm = e.get("provider_meta") or {}, e.get("sql_meta") or {}
        keep = e["step"] in ("routing", "branch") and e["action"] != "route" or pm or sm.get("query_ids") or e.get("engine_meta") or e["status"] == "engine_unavailable"
        if e["step"] == "handoff" and not last_handoff:
            keep = shown in (2, 9, 14)
        if not keep:
            continue
        if pm:
            extra, col = f"req={short(pm['requested_model'])} ans={short(pm['reported_model'])} [{pm['source']}]", BLUE
        elif sm.get("query_ids"):
            extra, col = f"{sm['source']}  query_id={sm['query_ids'][0]}", TEAL
        elif e.get("engine_meta") or e["status"] == "engine_unavailable":
            col = AMBER if e["status"] != "ok" else PURPLE
            extra = f"-> {e['status']}"
        else:
            extra, col = "", MUTED
        line(f"  {e['step']:<15} {e['agent']:<13} {e['action'][:24]:<25} {extra}", col)
        shown += 1
        if shown >= 17:
            break
    line("  ...", MUTED)
    line("$ python scripts/judge_compare.py   # all three examples match output_examples/  (also with --network none)", GREEN, True)
    return im


def pipeline_frame() -> Image.Image:
    im, d = canvas()
    header(d, "One question. Models propose. Code verifies.", "Bounded loop: 3 iterations, 24 model calls, 900 s")
    steps = [("Scout / routing", "validate + route", MUTED), ("Planner", "Nemotron Lightning", BLUE), ("Zetaris SQL", "retrieval is code", TEAL),
             ("Finder", "proposes links", BLUE), ("Verifier", "deterministic code", GREEN), ("Skeptic", "Nemotron Ultra", PURPLE), ("Reporter", "cites verified IDs", BLUE)]
    bw, gap, y0 = 214, 28, 330
    x = 100
    for i, (t, s, c) in enumerate(steps):
        box(d, [x, y0, x + bw, y0 + 170], fill=PANEL, outline=c)
        d.text((x + 16, y0 + 24), t, font=font(32, True), fill=TEXT)
        d.text((x + 16, y0 + 78), s, font=font(24), fill=c)
        if i < len(steps) - 1:
            d.text((x + bw + 3, y0 + 62), ">", font=font(40, True), fill=MUTED)
        x += bw + gap
    d.text((1000, 540), "Skeptic or verifier unhappy?  back to the Planner for more evidence", font=font(28), fill=AMBER)
    d.line([(1450, 520), (1450, 500), (580, 500), (580, 520)], fill=AMBER, width=4)
    pts = ["Duplicate reports never count as independent corroboration.",
           "Contradictions are surfaced, not merged away.",
           "\"No supportable link\" is a valid answer, with the reason.",
           "Every step, handoff, query id and model call lands in logs/trace-<run_id>.jsonl."]
    y = 650
    for p in pts:
        d.ellipse([110, y + 12, 122, y + 24], fill=TEAL)
        d.text((145, y), p, font=font(34), fill=TEXT)
        y += 62
    return im


def problem_frame() -> Image.Image:
    im, d = canvas()
    header(d, "Crime crosses station lines. Records don't.", "Three stations, five sources, almost no explicit links")
    cols = [("NORTH", BLUE, "Mon: office safe, rear door\nprised open, 36 mm\nchisel-tip tool marks"), ("CENTRAL", PURPLE, "Thu: office safe, rear door\nprised open, same\ntool-mark signature"),
            ("RIVERSIDE", GREEN, "Unrelated look-alike:\nsimilar wording, but\na forensic exclusion")]
    x = 100
    for name, c, txt in cols:
        box(d, [x, 280, x + 540, 640], fill=PANEL, outline=c)
        d.text((x + 28, 305), name + " STATION", font=font(36, True), fill=c)
        d.multiline_text((x + 28, 380), txt, font=font(30), fill=TEXT, spacing=10)
        d.text((x + 28, 570), "no link to the others", font=font(26), fill=MUTED)
        x += 590
    para(d, (100, 700), "Each station also holds its own CCTV, witness statements, forensic reports and press coverage. Evidence conflicts, and decoys look convincing. "
         "Someone has to read everything to see the pattern.", font(34), TEXT, width=1700)
    return im


def sponsors_frame() -> Image.Image:
    im, d = canvas()
    header(d, "Sponsor technology, used for real", "Each claim is something the trace can show")
    rows = [("Zetaris NDP", TEAL, "All retrieval is read-only SQL over 5 federated tables. Every query id is in the trace."),
            ("Meterless Nightshift", BLUE, "The only model path. Requested and answering model and receipt id recorded per call."),
            ("NVIDIA Nemotron", GREEN, "Lightning 30B: Planner, Finder, Reporter.  Ultra 550B: the Skeptic."),
            ("H-MEM + Markovian", PURPLE, "Run in the image from Meterless's open Apache-2.0 reference implementations."),
            ("World Model + Scout", AMBER, "Run when the Nightshift engine dir is mounted; otherwise traced engine_unavailable, never faked.")]
    y = 250
    for t, c, s in rows:
        box(d, [100, y, W - 100, y + 118], fill=PANEL, outline=c)
        d.text((130, y + 16), t, font=font(38, True), fill=c)
        para(d, (130, y + 68), s, font(28), TEXT, width=1640, gap=4)
        y += 138
    return im


def results_frame() -> Image.Image:
    im, d = canvas()
    header(d, "What it found, shown as it is", "Recorded Nemotron responses over the synthetic data (judge examples)")
    cards = [("Catalytic converters", "CENTRAL + RIVERSIDE", "3 WEAK links", "with the reasons a series can't be asserted", AMBER),
             ("Safe burglaries", "NORTH + CENTRAL", "no supportable link", "no independent corroboration retrieved", RED),
             ("Keyless SUV relay", "NORTH + RIVERSIDE", "no supportable link", "evidence too thin to assert a series", RED)]
    x = 100
    for t, s, r, why, c in cards:
        box(d, [x, 270, x + 540, 600], fill=PANEL, outline=c)
        d.text((x + 28, 295), t, font=font(38, True), fill=TEXT)
        d.text((x + 28, 350), s, font=font(26), fill=MUTED)
        d.text((x + 28, 430), r, font=font(40, True), fill=c)
        para(d, (x + 28, 500), why, font(26), TEXT, width=490)
        x += 590
    para(d, (100, 650), "The verifier's job is to refuse what the evidence does not support. We show that honestly rather than hide it. "
         "Leads, not verdicts: a human decides. All data is synthetic.", font(34), TEXT, width=1700)
    d.text((100, 860), "Reproduce: docker run -e SAMPLE_MODE=true oah-submission   |   scripts/judge_compare.py", font=mono(28), fill=TEAL)
    return im


def title_frame() -> Image.Image:
    im, d = canvas()
    d.text((100, 300), "AEGIS", font=font(150, True), fill=TEAL)
    d.text((100, 480), "Cross-Station Crime Intelligence", font=font(64, True), fill=TEXT)
    para(d, (100, 580), "Agents that link crimes across police stations, and show the evidence, the doubts and the missing pieces behind every link.", font(36), MUTED, width=1500)
    d.text((100, 820), "Track: Solving Fragmented Intelligence   ·   Open Agent Hackathon 2026", font=font(30), fill=MUTED)
    d.text((100, 880), "Zetaris NDP  ·  Meterless Nightshift & engines  ·  NVIDIA Nemotron", font=font(30), fill=TEAL)
    return im


def closing_frame() -> Image.Image:
    im, d = canvas()
    header(d, "Try it in two minutes", "No credentials, no network, no login")
    cmds = ["git clone <repo> && cd <repo>", "docker build -t oah-submission .", "docker run --rm -p 8000:8000 -e SAMPLE_MODE=true oah-submission",
            "curl -X POST localhost:8000/run -H 'Content-Type: application/json' -d @input_examples/example_1.json", "docker logs <container>   # one TRACE line per event"]
    box(d, [100, 260, W - 100, 640], fill=(10, 14, 24), outline=PANEL2)
    y = 290
    for c in cmds:
        d.text((130, y), "$ " + c, font=mono(28), fill=GREEN)
        y += 64
    para(d, (100, 700), "Leads, not verdicts. Synthetic data only. Models propose, code verifies, a Skeptic attacks.", font(40, True), TEXT, width=1700)
    d.text((100, 900), "Thank you.", font=font(48, True), fill=TEAL)
    return im


NARRATION = [
    ("title", "Crime doesn't stop at a station boundary. Records do. This is AEGIS, cross-station crime intelligence, built for the Solving Fragmented Intelligence track."),
    ("problem", "A burglar pries open an office safe at the North station on Monday, and another at Central on Thursday. Each station files one incident, and each holds its own camera footage, witness statements and forensic reports. Almost none of it is linked, so the pattern only shows to someone who reads everything. And real evidence conflicts, with look-alike decoys mixed in."),
    ("pipeline", "A detective asks one plain-language question. A Planner decides what to retrieve. Retrieval is plain SQL through Zetaris, never a model. A Finder proposes links. A deterministic verifier re-checks every ID and quote in code. A separate Skeptic model attacks the result, and can send the Planner back for more evidence. Only then does the Reporter write, and its citations are checked again."),
    ("overview", "This is the real application, filmed in offline sample mode, as the badge in the corner says. It runs over ten thousand synthetic records, across three stations and five sources. With credentials it reads the same data live from Zetaris."),
    ("query", "Here I ask whether the catalytic converter thefts in Central and Riverside could be one series."),
    ("graph", "The knowledge graph draws only what the result contains: incidents, the evidence they cite, and the links the Finder proposed. Nothing is invented to fill the picture."),
    ("result", "Each link comes back labelled by the verifier, here weak, with zero independent sources, and a score that is marked uncalibrated, not a probability. Contradictions, limitations and follow-up steps are listed. Nothing is presented as a verdict."),
    ("trace", "Every step is auditable: which agent ran, which Nemotron model answered, each query id, and each engine event. A step that made no model call shows no model and no receipt."),
    ("terminal", "For judges there is one command, with no credentials. The same agent loop runs offline, with local retrieval and recorded Nemotron responses replayed and labelled as replayed. The response lists the agents and engines, and the trace shows routing, handoffs, the Skeptic branch, query ids, engine events, and the requested and answering model on every call."),
    ("sponsors", "Zetaris is the data layer. Nightshift is the only model path. NVIDIA Nemotron does the reasoning: Lightning for planning, finding and reporting, and the much larger Ultra as the Skeptic. Meterless H-MEM and Markovian run inside the container from the open reference implementations. World Model and Scout run when the Nightshift engine directory is mounted, and are reported as unavailable otherwise, never faked."),
    ("results", "What did it find? With recorded Nemotron responses, the catalytic converter question returns three weak links, with the reasons a series can't be asserted. The safe burglary and keyless relay questions end in no supportable link, because no independent evidence corroborated them. We show that as it is. The verifier's job is to refuse what the evidence doesn't support."),
    ("closing", "Leads, not verdicts. All data is synthetic, and a human decides. Clone the repo, run one command, and read the trace yourself. Thank you."),
]


def build_video(frames: Path, out: Path, api_data: dict) -> Path:
    vf = out / "video_frames"
    vf.mkdir(parents=True, exist_ok=True)
    scenes = {
        "title": title_frame(), "problem": problem_frame(), "pipeline": pipeline_frame(),
        "overview": screen_frame("The dashboard, in offline sample mode", "Real app on local synthetic data (LOCAL FALLBACK badge): 3,000 incidents + 7,000 evidence records, 3 stations, 5 sources", shot(frames, "01_overview")),
        "query": screen_frame("Ask one plain-language question", "Investigation Workbench: Planner, Finder, Skeptic and Reporter agents", shot(frames, "05_investigation_query")),
        "graph": screen_frame("Knowledge graph of the result", "Only incidents, cited evidence and Finder links from the response are drawn; percentages are uncalibrated heuristics",
                              shot(frames, "09_scroll_1")),
        "result": screen_frame("Candidate links, with doubts and missing evidence", "Offline sample path: recorded Nemotron responses, replayed and labelled. Verifier status, sources, limitations, follow-ups.",
                               shot(frames, "09_scroll_4")),
        "trace": screen_frame("Execution trace", "Models, receipts and latencies appear only where the run recorded them (replayed Nemotron calls shown)", shot(frames, "10_trace_drawer")),
        "terminal": terminal_frame(api_data), "sponsors": sponsors_frame(), "results": results_frame(), "closing": closing_frame(),
    }
    clips = []
    for name, text in NARRATION:
        png, aiff, wav, mp4 = vf / f"{name}.png", vf / f"{name}.aiff", vf / f"{name}.wav", vf / f"{name}.mp4"
        scenes[name].save(png)
        subprocess.run(["say", "-v", "Samantha", "-r", "172", "-o", str(aiff), text], check=True)
        subprocess.run([FFMPEG, "-y", "-loglevel", "error", "-i", str(aiff), "-ar", "44100", "-ac", "1", str(wav)], check=True)
        with wave.open(str(wav)) as w:
            dur = w.getnframes() / w.getframerate()
        subprocess.run([FFMPEG, "-y", "-loglevel", "error", "-loop", "1", "-i", str(png), "-i", str(wav), "-t", f"{dur + 0.9:.2f}", "-af", "apad=pad_dur=1",
                        "-c:v", "libx264", "-tune", "stillimage", "-pix_fmt", "yuv420p", "-r", "24", "-c:a", "aac", "-b:a", "160k", "-shortest", str(mp4)], check=True)
        clips.append(mp4)
        print(f"scene {name}: {dur + 0.9:.1f}s")
    lst = vf / "list.txt"
    lst.write_text("".join(f"file '{c.name}'\n" for c in clips))
    final = out / "AEGIS_Demo_Video.mp4"
    subprocess.run([FFMPEG, "-y", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", str(lst), "-c", "copy", str(final)], check=True)
    (out / "narration.txt").write_text("\n\n".join(f"[{n}]\n{t}" for n, t in NARRATION) + "\n")
    return final


# ───────────────────────── slide deck ─────────────────────────
DARK = RGBColor(0x07, 0x0B, 0x14)
PANELC = RGBColor(0x11, 0x18, 0x27)
FG = RGBColor(0xE6, 0xEC, 0xF5)
MUT = RGBColor(0x8C, 0x98, 0xAA)
TEALC = RGBColor(0x2D, 0xD4, 0xBF)
AMBERC = RGBColor(0xF5, 0x9E, 0x0B)

SLIDES = [
    dict(title="AEGIS: Cross-Station Crime Intelligence", point="Title", sub="Agents that link crimes across police stations, and show the evidence, the doubts and the missing pieces behind every link.",
         bullets=["Track: Solving Fragmented Intelligence · Open Agent Hackathon 2026", "Zetaris NDP · Meterless Nightshift and engines · NVIDIA Nemotron"], img="01_overview",
         notes="Hook: crime doesn't stop at a station boundary, records do."),
    dict(title="The problem: each station sees only its own cases", point="Required point 1: problem and users",
         bullets=["The same tool-mark signature appears in NORTH on Monday and CENTRAL on Thursday. Two stations each file one burglary; nobody connects them.",
                  "Evidence is fragmented too: police incidents, CCTV, witness statements, forensic reports, press.",
                  "Most records carry no explicit link, so the link has to be deduced. Real evidence conflicts, and decoys look convincing.",
                  "Users: cross-station detectives and analysts (find and test links), supervisors (review), prosecutors (defensible record).",
                  "Our data is synthetic: 10,000 records, 5 sources, 3 stations, 8 planted relationships including 2 decoys."],
         notes="Say 'proposed lead', never 'caught'."),
    dict(title="The solution: one question, a checked answer", point="Required point 2: the solution",
         bullets=["Input: one plain-language question, e.g. 'Are the safe burglaries across NORTH and CENTRAL linked by the same pry bar?'",
                  "Agents read the sources through Zetaris and reason with Nemotron through Nightshift.",
                  "A deterministic verifier and a separate Skeptic check every link before a report is written.",
                  "Output: candidate links labelled SUPPORTED / WEAK / UNSUPPORTED with shared attributes, differences, missing evidence, objections and citations.",
                  "It can say no: without independent corroboration the answer is 'no supportable link', with the reason."], img="09_scroll_4"),
    dict(title="Agent workflow: models propose, code verifies, a Skeptic attacks", point="Required point 3: workflow and architecture",
         bullets=["Scout/routing, then (H-MEM recall), Planner, Zetaris SQL (code), Finder, Verifier (code), Skeptic, back to the Planner if more evidence is needed, Reporter, citation check, H-MEM store.",
                  "Bounded: 3 iterations, 24 model calls, 900 s by default; every Zetaris and Nightshift call has a timeout and a traced fallback.",
                  "Stopping reasons include evidence_sufficient, no_supportable_links_found, budget_exhausted, interrupted (resumable).",
                  "Not a fixed chain: routing, re-planning, adversarial critique, and code that has the last word."]),
    dict(title="Architecture and stack", point="Required points 3 and 4: architecture, tech stack",
         bullets=["Web: React 18, Vite, TypeScript: knowledge graph, execution-trace drawer, evidence pinboard.",
                  "API: FastAPI and Pydantic with typed contracts per agent (PlannerDecision, FinderOutput, SkepticReview, ReportOutput).",
                  "Data: Zetaris NDP, five federated tables. Models: NVIDIA Nemotron behind Meterless Nightshift.",
                  "Engines: Meterless H-MEM, Markovian (open references vendored, Apache-2.0), World Model and Scout (mounted).",
                  "Delivery: Docker image, POST /run and GET /health, JSONL traces, CI that runs the judge path and an offline check."]),
    dict(title="Zetaris: the load-bearing data layer", point="Required point 5: sponsor technology (Zetaris)",
         bullets=["Five federated tables: police_incidents (3,000), cctv (2,000), witness (2,000), forensic (1,500), newspaper (1,500).",
                  "One choke point for retrieval: read-only SELECT, bounded LIMIT, allow-listed characters; every query id is recorded in the trace.",
                  "The model never touches the database: it sees at most 14 incidents and 8 records per source, and every ID it cites is re-checked.",
                  "Live run from inside Docker (reported by the team; trace not committed): data_mode zetaris_live, query UUIDs, all five tables.",
                  "Honest fallback: with POST /run an unreachable Zetaris falls back to local sample data and the trace says so."]),
    dict(title="Meterless engines and Nightshift", point="Required point 5: sponsor technology (Meterless)",
         bullets=["World Model: shared ledger of every decision, finding revision, objection and verdict. (test_meterless_engine.py)",
                  "H-MEM: recalls vetted leads as context only; stores only verified, Skeptic-reviewed leads (0.6 floor, secret gate). (test_engine_adapters.py, test_reference_engines.py)",
                  "Markovian: validated chunk config, 800-token carry-over, resumable runs. (test_engine_adapters.py, test_reference_engines.py)",
                  "Scout Intent: detector on every request as a safety signal; deterministic routing; refuses destructive requests. (test_engine_adapters.py)",
                  "Nightshift: the only model path; requested and answering model, receipt, latency, tokens on every call. An engine that cannot run is traced engine_unavailable, never faked."], img="10_trace_drawer"),
    dict(title="NVIDIA Nemotron: where it contributes", point="Required point 5: sponsor technology (NVIDIA)",
         bullets=["Lightning 30B-A3B: Planner, Finder, Reporter (fast lane).", "Ultra 550B-A55B: Skeptic, a different model from the Finder and never shown its score (deep lane, about 20 to 40 s per call).",
                  "Every response names the model that answered, taken from the response, not from config; fallback to Lightning is explicit with the reason.",
                  "Not used, and we say so: Embed 1B is listed but /v1/embeddings returned 404; Rerank and Bolt are not served; no vision claim."]),
    dict(title="What is unique", point="Required point 6: what is unique",
         bullets=["1. 'Models propose. Code verifies. A Skeptic attacks.' IDs, verbatim quotes, shared attributes on both incidents, station and timeline are re-checked in code.",
                  "2. 'Duplicates aren't corroboration.' A repeated camera window, one witness reported twice, a newspaper article about a police-recorded incident never count as a second source.",
                  "3. 'Memory only learns what survived scrutiny.' H-MEM stores only verified, Skeptic-reviewed leads; runs checkpoint and resume.",
                  "4. Runs anywhere: the real loop offline in one Docker command, and the same image goes live."]),
    dict(title="Product visuals", point="Required point 7: visuals", bullets=["Live dashboard, investigation workbench, knowledge graph of the result, link cards with verifier status, execution trace. All data synthetic."], img="09_scroll_1"),
    dict(title="Results, shown honestly", point="Results",
         bullets=["Judge examples (recorded Nemotron responses, offline sample path): catalytic converters give 3 WEAK links with reasons a series can't be asserted; safe burglaries and keyless relay end in 'no supportable link'.",
                  "Reported by the team (trace not committed): live run through Docker with real Zetaris and Nightshift: full loop, real Nemotron on all four agents, zero partial failures. A live safe-burglary run retrieved the planted pair but the verifier and Skeptic refused an ungrounded link.",
                  "We do not claim a model benchmark. evaluation/results.json used a scripted offline stand-in for the models, so it tests the pipeline and scoring mechanics, not model quality.",
                  "Engineering: 201 test functions; one command (pytest backend/tests tests) passes, engine and live tests skip when unavailable; secret scan on source and image; CI runs the judge path and --network none."]),
    dict(title="Responsible AI", point="Responsible AI",
         bullets=["Leads, not verdicts: no determination of guilt; a human decides.", "Explainable by construction: shared attributes, differences, missing evidence, objections, citations, query ids.",
                  "Models cannot invent evidence: every ID, quote and attribute is re-checked in code.", "Scores are labelled uncalibrated heuristics.",
                  "Transparent degradation: replayed, code-fallback and local-sample steps are labelled in the trace.", "Synthetic data only; no secrets in the repo or image; known bias risk if used on real records."]),
    dict(title="Future enhancements and limitations", point="Required point 8: future and limitations",
         bullets=["Next: semantic retrieval and reranking once the routes are served; CCTV vision; crime-specific Scout intents; streamed progress; bound SQL parameters; calibrated confidence; larger multi-run benchmark with real Nemotron; analyst feedback.",
                  "Limits: synthetic data; only H-MEM and Markovian run in the image; World Model and Scout need a mounted Nightshift engine directory; verifier heuristics cover known patterns; Ultra is slow; sample-mode quality depends on recorded fixtures.",
                  "Real-world use would need legal basis, consent, oversight, and a privacy and bias review."]),
    dict(title="Run it yourself", point="Judge path", bullets=["docker build -t oah-submission .", "docker run --rm -p 8000:8000 -e SAMPLE_MODE=true oah-submission",
                                                                "curl -X POST localhost:8000/run -H 'Content-Type: application/json' -d @input_examples/example_1.json",
                                                                "docker logs <container>   (one TRACE line per event)", "No credentials, no network, no login."]),
]


def fit_font(bullets: list[str], width_in: float, height_in: float) -> int:
    """Largest font size (pt) at which the bullets are estimated to fit the box (Calibri-ish: ~0.5 em per char)."""
    import math
    for pt in (17, 16, 15, 14, 13, 12, 11, 10):
        cpl = max(10, int(width_in * 72 / (pt * 0.52)))
        h = sum(math.ceil(len(b) / cpl) * pt * 1.22 + 9 for b in bullets) / 72
        if h <= height_in:
            return pt
    return 10


def style(par, size: int, color, bold: bool = False) -> None:
    """Set the font on the paragraph AND its runs (PowerPoint honours run-level properties)."""
    par.font.size, par.font.color.rgb, par.font.bold = Pt(size), color, bold
    for r in par.runs:
        r.font.size, r.font.color.rgb, r.font.bold = Pt(size), color, bold


def build_deck(frames: Path, out: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for i, s in enumerate(SLIDES):
        sl = prs.slides.add_slide(blank)
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = DARK
        bar = sl.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.08))
        bar.fill.solid(); bar.fill.fore_color.rgb = TEALC; bar.line.fill.background()
        tb = sl.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0)).text_frame
        tb.word_wrap = True
        p = tb.paragraphs[0]; p.text = s["title"]; style(p, 32 if len(s["title"]) < 50 else 28, FG, True)
        if s.get("sub"):
            q = tb.add_paragraph(); q.text = s["sub"]; style(q, 16, MUT)
        img = shot(frames, s["img"]) if s.get("img") else None
        width = Inches(6.6) if img is not None else Inches(12.1)
        bx = sl.shapes.add_textbox(Inches(0.6), Inches(1.7), width, Inches(5.2)).text_frame
        bx.word_wrap = True
        for j, b in enumerate(s["bullets"]):
            para_ = bx.paragraphs[0] if j == 0 else bx.add_paragraph()
            para_.text = b if not s.get("sub") else b
            style(para_, fit_font(s["bullets"], 6.5 if img is not None else 12.0, 5.0), FG); para_.space_after = Pt(9)
        if img is not None:
            tmp = out / "video_frames" / f"_slide_{i}.png"
            img.save(tmp)
            sl.shapes.add_picture(str(tmp), Inches(7.4), Inches(1.7), width=Inches(5.5))
        ft = sl.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.35)).text_frame
        ft.paragraphs[0].text = f"{s['point']}   ·   {i + 1}/{len(SLIDES)}"
        style(ft.paragraphs[0], 11, MUT)
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]
    dest = out / "AEGIS_Slide_Deck.pptx"
    prs.save(dest)
    return dest


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--frames", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--api", default="http://127.0.0.1:8000")
    ap.add_argument("--skip-video", action="store_true")
    ap.add_argument("--container", help="docker container serving --api; the terminal scene then shows the real judge path")
    a = ap.parse_args()
    a.out.mkdir(parents=True, exist_ok=True)
    (a.out / "video_frames").mkdir(exist_ok=True)
    api_data = load_api(a.api, container=a.container)
    print("deck:", build_deck(a.frames, a.out))
    if not a.skip_video:
        print("video:", build_video(a.frames, a.out, api_data))
    return 0


if __name__ == "__main__":
    sys.exit(main())
